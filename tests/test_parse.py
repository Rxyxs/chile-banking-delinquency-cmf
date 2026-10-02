"""El parser se prueba con archivos sintéticos que reproducen los dos formatos de la CMF."""
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
import pytest

from cmf_delinquency.parse import (
    SYSTEM_NAME,
    build_panel,
    canonical_bank,
    detect_layout,
    parse_file,
)

BANKS = [
    ("Banco Uno", [1.0, 1.1, 0.9, 1.2, 1.3, 1.0], 1000.0),
    ("Banco Dos", [2.0, 2.1, 1.9, "---", "---", "---"], 500.0),
    (SYSTEM_NAME, [1.3, 1.4, 1.2, 1.2, 1.3, 1.0], 1500.0),
]


def _write_old(path: Path, banks=BANKS) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["INDICADORES DE RIESGO"] + [None] * 8)
    ws.append(["Instituciones", None, None, None, None, None, None, None, "Cartera morosa (MM$)"])
    ws.append([None, "Colocaciones", "Total", "Comerciales (Empresas)", "Personas", None, None, None, None])
    ws.append([None, None, None, None, "Total", "Consumo", "Vivienda", None, None])
    for name, vals, mm in banks:
        ws.append([name, *vals, None, mm])
    ws.append(["Notas:"] + [None] * 8)
    wb.save(path)


def _write_new(path: Path, banks=BANKS) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([None] * 12)
    ws.append([None, "Instituciones (*)", "Cartera con morosidad de 90 días o más (%)"] + [None] * 7 + ["MM$", None])
    ws.append([None, None, "Colocaciones (costo amortizado y valor razonable)", "Colocaciones a costo amortizado"] + [None] * 8)
    ws.append([None, None, None, "Total", "Comerciales", "Personas", None, None, "Adeudado por bancos", None, None, None])
    ws.append([None, None, None, None, None, "Total", "Consumo", "Vivienda", None, None, None, None])
    for name, vals, mm in banks:
        ws.append([None, name, *vals, "---", None, mm, mm])
    ws.append([None, "Notas:"] + [None] * 10)
    wb.save(path)


def test_parse_old_layout(tmp_path):
    f = tmp_path / "2019-05.xlsx"
    _write_old(f)
    out = parse_file(f)
    assert list(out["bank_raw"]) == ["Banco Uno", "Banco Dos", SYSTEM_NAME]
    assert out["layout"].eq("antiguo").all()
    uno = out.iloc[0]
    assert (uno["total"], uno["clientes"], uno["comercial"]) == (1.0, 1.1, 0.9)
    assert (uno["personas"], uno["consumo"], uno["vivienda"]) == (1.2, 1.3, 1.0)
    assert uno["mora90_mm"] == 1000.0
    assert out["date"].iloc[0] == pd.Timestamp("2019-05-01")


def test_parse_new_layout(tmp_path):
    f = tmp_path / "2024-02.xlsx"
    _write_new(f)
    out = parse_file(f)
    assert out["layout"].eq("nuevo").all()
    uno = out.iloc[0]
    assert (uno["total"], uno["comercial"], uno["consumo"], uno["vivienda"]) == (1.0, 0.9, 1.3, 1.0)
    assert uno["mora90_mm"] == 1000.0


def test_dashes_become_nan_not_zero(tmp_path):
    f = tmp_path / "2024-02.xlsx"
    _write_new(f)
    dos = parse_file(f).set_index("bank_raw").loc["Banco Dos"]
    assert np.isnan(dos["personas"]) and np.isnan(dos["consumo"]) and np.isnan(dos["vivienda"])
    assert dos["comercial"] == 1.9


def test_system_row_stops_the_parse(tmp_path):
    """Lo que viene después de 'Sistema Bancario' son notas al pie y no se lee."""
    f = tmp_path / "2024-02.xlsx"
    _write_new(f)
    assert parse_file(f)["bank_raw"].iloc[-1] == SYSTEM_NAME


def test_missing_system_row_raises(tmp_path):
    f = tmp_path / "2024-02.xlsx"
    _write_new(f, banks=BANKS[:2])
    with pytest.raises(ValueError, match="Sistema Bancario"):
        parse_file(f)


def test_unrecognised_header_raises(tmp_path):
    f = tmp_path / "2019-05.xlsx"
    _write_old(f)
    wb = openpyxl.load_workbook(f)
    for row in wb.active.iter_rows(min_row=1, max_row=4):
        for cell in row:
            if isinstance(cell.value, str) and cell.value in {"Consumo", "Vivienda", "Comerciales (Empresas)"}:
                cell.value = "otra cosa"
    wb.save(f)
    with pytest.raises(ValueError, match="Formato no reconocido"):
        detect_layout(pd.read_excel(f, header=None))


def test_canonical_names():
    d = pd.Timestamp("2020-01-01")
    assert canonical_bank("Itaú Corpbanca", d) == "Banco Itaú Chile"
    assert canonical_bank("Jp Morgan Chase Bank, N.A.", d) == "JP Morgan Chase Bank, N.A."
    assert canonical_bank("Banco de Chile", d) == "Banco de Chile"
    assert canonical_bank("Scotiabank Azul", d) == canonical_bank("Banco Bilbao Vizcaya Argentaria, Chile", d)


def test_pre_merger_itau_is_not_the_same_series():
    assert canonical_bank("Banco Itaú Chile", pd.Timestamp("2016-03-01")) is None
    assert canonical_bank("Banco Itaú Chile", pd.Timestamp("2016-04-01")) == "Banco Itaú Chile"
    assert canonical_bank("Banco Itaú Chile", pd.Timestamp("2024-01-01")) == "Banco Itaú Chile"


def test_build_panel_mixes_layouts(tmp_path):
    _write_old(tmp_path / "2021-12.xlsx")
    _write_new(tmp_path / "2022-01.xlsx")
    panel = build_panel(tmp_path)
    assert sorted(panel["layout"].unique()) == ["antiguo", "nuevo"]
    assert panel["is_system"].sum() == 2
    assert not panel.duplicated(["date", "bank"]).any()


def test_build_panel_without_files_fails(tmp_path):
    with pytest.raises(FileNotFoundError):
        build_panel(tmp_path)
